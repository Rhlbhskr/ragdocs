import io
import zipfile
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from app.main import app
from app.parsing import chunk_pages, parse_document
from app.grounding import answer_question


@pytest.fixture(scope='module')
def client():
    with TestClient(app) as c:
        assert c.post('/api/workspace').status_code == 200
        yield c


def test_upload_question_citations_history_and_source(client):
    result = client.post('/api/documents', files={'file': ('Orchid.txt', b'Orchid project launch date is November 17, 2026. The project budget is 42000 dollars.')})
    assert result.status_code == 201, result.text
    identifier = result.json()['id']
    result = client.post('/api/chat', json={'question': 'What is the Orchid project launch date?', 'mode': 'hybrid', 'top_k': 3})
    assert result.status_code == 200, result.text
    turn = result.json()
    assert 'November 17, 2026' in turn['answer']
    assert turn['citations'] and any(c['document_id'] == identifier for c in turn['citations'])
    assert all(c['text'] for c in turn['citations'])
    assert client.get('/api/history').json()[-1]['question'] == turn['question']
    assert client.get(f'/api/documents/{identifier}/source').json()['passages'][0]['page'] is None
    assert client.delete(f'/api/documents/{identifier}').status_code == 204
    assert client.get(f'/api/documents/{identifier}/source').status_code == 404
    assert client.get('/api/history').json() == []


def test_workspace_isolation(client):
    result = client.post('/api/documents', files={'file': ('Secret.txt', b'Private quasar password is violet.')})
    identifier = result.json()['id']
    cookie = client.cookies.get('grounded_workspace')
    client.cookies.clear()
    assert client.get('/api/documents').status_code == 401
    client.post('/api/workspace')
    assert identifier not in [d['id'] for d in client.get('/api/documents').json()]
    assert client.get(f'/api/documents/{identifier}/source').status_code == 404
    assert client.delete(f'/api/documents/{identifier}').status_code == 404
    turn = client.post('/api/chat', json={'question': 'What is the private quasar password?'}).json()
    assert all(p['document_id'] != identifier for p in turn['passages'])
    client.cookies.clear()
    client.cookies.set('grounded_workspace', cookie)
    client.delete(f'/api/documents/{identifier}')


def test_low_evidence_abstention(client):
    client.delete('/api/history')
    result = client.post('/api/chat', json={'question': 'Explain quantum entanglement experiments.'}).json()
    assert result['generation'] == 'abstention'
    assert result['citations'] == []
    assert not result['grounding']['supported']


def test_invalid_uploads_and_inputs(client):
    assert client.post('/api/documents', files={'file': ('bad.exe', b'hello')}).status_code == 422
    assert client.post('/api/documents', files={'file': ('bad.pdf', b'not a pdf')}).status_code == 422
    assert client.post('/api/documents', files={'file': ('empty.txt', b'  ')}).status_code == 422
    assert client.post('/api/documents', files={'file': ('test.txt', b'hello')}, data={'chunk_size': 1}).status_code == 422
    assert client.post('/api/chat', json={'question': '  '}).status_code == 422
    assert client.post('/api/chat', json={'question': 'test', 'top_k': 99}).status_code == 422


def test_docx_paragraphs_and_tables():
    doc = Document(); doc.add_paragraph('Policy text')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Allowance'; table.cell(0, 1).text = '$750'
    stream = io.BytesIO(); doc.save(stream)
    pages = parse_document('policy.docx', stream.getvalue())
    assert pages[0][0] is None
    assert 'Policy text' in pages[0][1] and '$750' in pages[0][1]


def test_scanned_pdf_rejected():
    writer = PdfWriter(); writer.add_blank_page(width=600, height=800)
    stream = io.BytesIO(); writer.write(stream)
    with pytest.raises(ValueError, match='OCR'):
        parse_document('scan.pdf', stream.getvalue())


def test_pdf_upload_with_page_citations(client):
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in ('A general introduction.', 'The zephyr spacecraft departure time is 06:45 UTC.'):
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(f'BT /F1 12 Tf 50 750 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output)
    upload = client.post('/api/documents', files={'file': ('Flight.pdf', output.getvalue())})
    assert upload.status_code == 201
    client.delete('/api/history')
    result = client.post('/api/chat', json={'question': 'What is the zephyr spacecraft departure time?'}).json()
    citation = next(c for c in result['citations'] if c['document_id'] == upload.json()['id'])
    assert citation['page'] == 2 and '06:45' in result['answer']
    client.delete('/api/documents/' + upload.json()['id'])


def test_request_body_limit(client):
    response = client.post('/api/documents', content=b'x' * (16 * 1024 * 1024 + 1), headers={'Content-Type': 'application/octet-stream'})
    assert response.status_code == 413


def test_chunks_preserve_pages_and_overlap():
    pages = [(4, ' '.join(f'w{i}' for i in range(600))), (9, 'Policy number nine')]
    chunks = chunk_pages(pages, 256, 64)
    assert chunks[0]['page'] == 4 and chunks[-1]['page'] == 9
    assert chunks[0]['text'].split()[-64:] == chunks[1]['text'].split()[:64]
    assert len(chunks[0]['text'].split()) == 256
    assert chunks[-1]['number'] == len(chunks)


def test_llm_invalid_citations_abstain(monkeypatch):
    from app.config import settings
    import app.grounding as grounding
    monkeypatch.setattr(settings, 'llm_base_url', 'https://example.invalid/v1')
    class FakeResponse:
        def raise_for_status(self): pass
        def json(self): return {'choices': [{'message': {'content': 'The budget is $750 [99].'}}]}
    monkeypatch.setattr(grounding.httpx.Client, 'post', lambda *a, **k: FakeResponse())
    result = answer_question('What is the budget?', [{'text': 'The budget is $750.', 'name': 'test', 'page': None, 'number': 1}], [])
    assert result['generation'] == 'abstention' and not result['citations']


def test_web_and_openapi(client):
    assert client.get('/').status_code == 200
    assert client.get('/static/app.js').status_code == 200
    assert '/api/chat' in client.get('/openapi.json').json()['paths']
    assert client.get('/api/health').json()['embedding_backend'] == 'hash'
