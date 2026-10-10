import io
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def clean(text: str) -> str:
    return re.sub(r'\s+', ' ', text.replace('\x00', '')).strip()


def parse_document(name: str, content: bytes, max_pages=500) -> list[tuple[int | None, str]]:
    suffix = Path(name).suffix.lower()
    if suffix == '.pdf':
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are not supported.')
        if len(reader.pages) > max_pages:
            raise ValueError(f'Documents must contain at most {max_pages} pages.')
        pages = [(i + 1, clean(page.extract_text() or '')) for i, page in enumerate(reader.pages)]
    elif suffix == '.docx':
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Expanded DOCX is too large.')
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        parts += [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        pages = [(None, clean('\n'.join(parts)))]
    elif suffix == '.txt':
        pages = [(None, clean(content.decode('utf-8-sig')))]
    else:
        raise ValueError('Upload a PDF, DOCX, or UTF-8 TXT file.')
    pages = [(page, text) for page, text in pages if text]
    if not pages:
        raise ValueError('No readable text found. Scanned PDFs need OCR before upload.')
    if sum(len(text) for _, text in pages) > 2_000_000:
        raise ValueError('Extracted text exceeds the 2 million character limit.')
    return pages


def chunk_pages(pages, size=512, overlap=64):
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError('Overlap must be smaller than chunk size.')
    chunks = []
    for page, text in pages:
        words = text.split()
        for start in range(0, len(words), size - overlap):
            chunks.append({'page': page, 'text': ' '.join(words[start:start + size]), 'number': len(chunks) + 1})
            if start + size >= len(words):
                break
    return chunks
